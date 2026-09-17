import io
import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Đặt màu nền cho cell trong Word"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    """Đặt lề trong cho cell (khoảng đệm chữ)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_packing_word_doc(shop_name="GIMME TEE", *args, **kwargs):
    """
    Tạo tài liệu Word (.docx) đúng chuẩn 100% như mẫu:
    - Tiêu đề: GIMME - SHOPEE - SÁNG - 17.09 - 418 ĐƠN - 409 ÁO
    - Bảng 4 cột: SKU sản phẩm | Màu | Size | Áo
    - Gom nhóm từng SKU và có dòng 'Tổng số {SKU}' màu xanh nhạt
    - Dòng 'Tổng cộng' ở cuối bảng
    """
    platform = kwargs.get('platform', 'TIKTOK')
    shift_name = kwargs.get('shift_name', 'SÁNG')
    work_date = kwargs.get('work_date', '')
    metrics = kwargs.get('metrics', {})
    pick_list_df = kwargs.get('pick_list_df', None)
    df_valid = kwargs.get('df_valid', None)

    if len(args) == 6:
        platform, shift_name, work_date, metrics, pick_list_df, df_valid = args
    elif len(args) == 5:
        shift_name, work_date, metrics, pick_list_df, df_valid = args
    elif len(args) >= 1:
        pos_names = ['platform', 'shift_name', 'work_date', 'metrics', 'pick_list_df', 'df_valid']
        for i, val in enumerate(args):
            if i < len(pos_names) and pos_names[i] not in kwargs:
                if pos_names[i] == 'platform': platform = val
                elif pos_names[i] == 'shift_name': shift_name = val
                elif pos_names[i] == 'work_date': work_date = val
                elif pos_names[i] == 'metrics': metrics = val
                elif pos_names[i] == 'pick_list_df': pick_list_df = val
                elif pos_names[i] == 'df_valid': df_valid = val

    doc = Document()
    
    # Thiết lập lề trang
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # 1. Format tên gian hàng và thông tin tiêu đề
    shop_tag = "GIMME" if "GIMME" in str(shop_name).upper() else str(shop_name).upper()
    plat_tag = "SHOPEE" if "SHOPEE" in str(platform).upper() else "TIKTOK"
    shift_tag = str(shift_name).upper()
    
    # Định dạng ngày: DD.MM (ví dụ 17.09)
    date_parts = str(work_date).split("/")
    if len(date_parts) >= 2:
        date_tag = f"{date_parts[0]}.{date_parts[1]}"
    else:
        date_tag = str(work_date)

    orders_cnt = metrics.get('valid_orders_count', 0)
    items_cnt = metrics.get('total_valid_items', 0)

    # Tiêu đề tài liệu chính xác theo mẫu: GIMME - SHOPEE - SÁNG - 17.09 - 418 ĐƠN - 409 ÁO
    header_title = f"{shop_tag} - {plat_tag} - {shift_tag} - {date_tag} - {orders_cnt} ĐƠN - {items_cnt} ÁO"

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    
    run_title = p_title.add_run(header_title)
    run_title.font.name = "Arial"
    run_title.font.size = Pt(12)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)

    # 2. Tạo bảng 4 cột: SKU sản phẩm | Màu | Size | Áo
    cols_width = [Inches(1.7), Inches(2.7), Inches(1.4), Inches(1.0)]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Thêm viền đen mảnh chuẩn Word
    tblPr = table._tbl.tblPr
    border_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(border_xml))

    # Header Row
    headers = ["SKU sản phẩm", "Màu", "Size", "Áo"]
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = cols_width[i]
        set_cell_background(hdr_cells[i], "D9E1F2") # Màu xám xanh nhạt chuẩn mẫu
        set_cell_margins(hdr_cells[i], top=60, bottom=60, left=80, right=80)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i == 3 else WD_ALIGN_PARAGRAPH.LEFT
        
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Nếu có đơn trống hoặc dòng trống, có thể thêm 'Tổng số 0'
    # Nhóm dữ liệu theo 'SKU sản phẩm'
    if not pick_list_df.empty:
        sku_groups = pick_list_df.groupby('SKU sản phẩm', sort=False)

        for sku, group in sku_groups:
            sku_subtotal = 0
            
            # Các dòng chi tiết từng Màu & Size trong SKU này
            for _, row in group.iterrows():
                r_cells = table.add_row().cells
                for c_i in range(4):
                    r_cells[c_i].width = cols_width[c_i]
                    set_cell_margins(r_cells[c_i], top=40, bottom=40, left=80, right=80)
                    r_cells[c_i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                # Cột 0: SKU sản phẩm
                p0 = r_cells[0].paragraphs[0]
                p0.paragraph_format.space_before = Pt(1)
                p0.paragraph_format.space_after = Pt(1)
                p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r0 = p0.add_run(str(sku))
                r0.font.name = "Arial"
                r0.font.size = Pt(9.5)

                # Cột 1: Màu
                p1 = r_cells[1].paragraphs[0]
                p1.paragraph_format.space_before = Pt(1)
                p1.paragraph_format.space_after = Pt(1)
                p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r1 = p1.add_run(str(row.get('Màu', '')))
                r1.font.name = "Arial"
                r1.font.size = Pt(9.5)

                # Cột 2: Size
                p2 = r_cells[2].paragraphs[0]
                p2.paragraph_format.space_before = Pt(1)
                p2.paragraph_format.space_after = Pt(1)
                p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r2 = p2.add_run(str(row.get('Size', '')))
                r2.font.name = "Arial"
                r2.font.size = Pt(9.5)

                # Cột 3: Áo (Số lượng)
                qty = int(row.get('Áo', 1))
                sku_subtotal += qty
                p3 = r_cells[3].paragraphs[0]
                p3.paragraph_format.space_before = Pt(1)
                p3.paragraph_format.space_after = Pt(1)
                p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r3 = p3.add_run(str(qty))
                r3.font.name = "Arial"
                r3.font.size = Pt(9.5)

            # Dòng Subtotal: Tổng số {SKU}
            sub_cells = table.add_row().cells
            for c_i in range(4):
                sub_cells[c_i].width = cols_width[c_i]
                set_cell_background(sub_cells[c_i], "D9E1F2") # Màu nền xám xanh nhạt
                set_cell_margins(sub_cells[c_i], top=50, bottom=50, left=80, right=80)
                sub_cells[c_i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            # Col 0: Tổng số {SKU}
            p_sub0 = sub_cells[0].paragraphs[0]
            p_sub0.paragraph_format.space_before = Pt(1)
            p_sub0.paragraph_format.space_after = Pt(1)
            p_sub0.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_sub0 = p_sub0.add_run(f"Tổng số {sku}")
            r_sub0.font.name = "Arial"
            r_sub0.font.size = Pt(9.5)
            r_sub0.font.bold = True
            r_sub0.font.color.rgb = RGBColor(0, 0, 0)

            # Col 3: Số lượng tổng của SKU này
            p_sub3 = sub_cells[3].paragraphs[0]
            p_sub3.paragraph_format.space_before = Pt(1)
            p_sub3.paragraph_format.space_after = Pt(1)
            p_sub3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_sub3 = p_sub3.add_run(str(sku_subtotal))
            r_sub3.font.name = "Arial"
            r_sub3.font.size = Pt(9.5)
            r_sub3.font.bold = True
            r_sub3.font.color.rgb = RGBColor(31, 78, 121) # Màu xanh đậm #1F4E79 như ảnh mẫu

        # Dòng Tổng cộng cuối bảng
        tot_cells = table.add_row().cells
        for c_i in range(4):
            tot_cells[c_i].width = cols_width[c_i]
            set_cell_background(tot_cells[c_i], "D9E1F2")
            set_cell_margins(tot_cells[c_i], top=60, bottom=60, left=80, right=80)
            tot_cells[c_i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        p_tot0 = tot_cells[0].paragraphs[0]
        p_tot0.paragraph_format.space_before = Pt(2)
        p_tot0.paragraph_format.space_after = Pt(2)
        r_tot0 = p_tot0.add_run("Tổng cộng")
        r_tot0.font.name = "Arial"
        r_tot0.font.size = Pt(10)
        r_tot0.font.bold = True

        p_tot3 = tot_cells[3].paragraphs[0]
        p_tot3.paragraph_format.space_before = Pt(2)
        p_tot3.paragraph_format.space_after = Pt(2)
        p_tot3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_tot3 = p_tot3.add_run(str(items_cnt))
        r_tot3.font.name = "Arial"
        r_tot3.font.size = Pt(10)
        r_tot3.font.bold = True
        r_tot3.font.color.rgb = RGBColor(31, 78, 121)

    # Lưu tài liệu ra bộ nhớ BytesIO
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
